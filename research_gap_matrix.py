"""Batch analysis PDF to Excel research-gap matrix using local Ollama."""
import hashlib, json, re
from pathlib import Path
import pandas as pd
from pypdf import PdfReader
from langchain_ollama import ChatOllama
ROOT=Path(__file__).resolve().parent; DATA_PATH=ROOT/'data'; OUTPUT_PATH=ROOT/'research_gap_matrix.xlsx'; MODEL='mistral'
FIELDS=['title','authors_year','doi','problem','method','dataset','findings','limitations','future_work','research_gap_candidate','evidence_or_pages']
UNKNOWN='Tidak disebutkan dalam teks yang tersedia'
PROMPT='''Analisis artikel akademik ini. Keluarkan HANYA JSON valid dengan keys: title, authors_year, doi, problem, method, dataset, findings, limitations, future_work, research_gap_candidate, evidence_or_pages. Semua nilai string. Jangan mengarang. Jika bukti tidak ada, tulis "Tidak disebutkan dalam teks yang tersedia". Bedakan fakta eksplisit dan interpretasi. Kandidat gap bukan kesimpulan final. Cantumkan halaman/bukti bila tersedia.\nNAMA FILE: {filename}\nTEKS ARTIKEL:\n{text}'''
def extract_pdf_text(path,max_chars=30000):
    reader=PdfReader(str(path)); pages=[]; total=0
    for i,page in enumerate(reader.pages,1):
        text=(page.extract_text() or '').strip()
        if text:
            piece=f'[Halaman {i}]\n{text}'; pages.append(piece); total+=len(piece)
        if total>=max_chars: break
    return '\n\n'.join(pages)[:max_chars]
def parse_json(raw):
    try: return json.loads(raw)
    except json.JSONDecodeError:
        match=re.search(r'\{.*\}',raw,re.S)
        if not match: raise
        return json.loads(match.group(0))
def main():
    pdfs=sorted(DATA_PATH.glob('*.pdf'))
    if not pdfs: raise SystemExit(f'Tidak ada PDF di {DATA_PATH}.')
    seen={}; unique=[]; duplicates=[]
    for pdf in pdfs:
        digest=hashlib.sha256(pdf.read_bytes()).hexdigest()
        if digest in seen: duplicates.append((pdf.name,seen[digest]))
        else: seen[digest]=pdf.name; unique.append(pdf)
    for duplicate,original in duplicates: print(f'Duplikat identik dilewati: {duplicate} = {original}')
    llm=ChatOllama(model=MODEL,temperature=0,format='json'); rows=[]
    for i,pdf in enumerate(unique,1):
        print(f'[{i}/{len(unique)}] Menganalisis {pdf.name}')
        try:
            text=extract_pdf_text(pdf)
            if not text.strip(): raise ValueError('Teks PDF kosong; mungkin PDF scan dan perlu OCR.')
            data=parse_json(llm.invoke(PROMPT.format(filename=pdf.name,text=text)).content)
            row={f:str(data.get(f,UNKNOWN)) for f in FIELDS}; row.update(source_file=pdf.name,sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),status='Perlu verifikasi manusia',error='')
        except Exception as exc:
            row={f:'' for f in FIELDS}; row.update(source_file=pdf.name,sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),status='Gagal dianalisis',error=str(exc)); print('  GAGAL:',exc)
        rows.append(row)
    cols=['source_file','sha256']+FIELDS+['status','error']; df=pd.DataFrame(rows,columns=cols)
    with pd.ExcelWriter(OUTPUT_PATH,engine='openpyxl') as writer:
        df.to_excel(writer,index=False,sheet_name='Research Gap'); ws=writer.sheets['Research Gap']; ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        from openpyxl.styles import Alignment,Font,PatternFill
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
        for cells in ws.columns:
            letter=cells[0].column_letter; ws.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in cells[:30])+2,14),42)
            for cell in cells[1:]: cell.alignment=Alignment(vertical='top',wrap_text=True)
    print(f'Selesai: {OUTPUT_PATH}\nVerifikasi semua klaim terhadap artikel asli sebelum digunakan dalam tesis.')
if __name__=='__main__': main()
